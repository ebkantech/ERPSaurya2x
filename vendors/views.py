from collections import defaultdict
from datetime import timedelta
from io import BytesIO
from pathlib import Path
import json
import sys

from django.contrib import messages
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Count, Q
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from accounts.forms import StaffProfileForm
from accounts.models import StaffProfile
from administration.models import SystemAuditLog
from administration.services import log_system_audit
from audit_logs.models import VendorActivityLog
from audit_logs.services import log_vendor_activity
from core.models import Vendor
from notifications.services import create_notification
from permissions.utils import (
    can_create_vendors,
    ensure_vendor_access,
    ensure_vendor_write_access,
    get_accessible_vendor_queryset,
    get_active_vendor_assignments,
    get_staff_profile,
    is_admin_like,
    is_view_only,
    require_admin_access,
    require_authenticated,
    role_label,
)
from purchase_orders.models import PurchaseOrder
from tasks.models import VendorTask

from .forms import (
    VendorAssignmentForm,
    VendorAssignmentUploadForm,
    VendorAutoDistributeForm,
    VendorBulkAssignForm,
    VendorMasterForm,
    VendorNoteForm,
)
from .models import VendorAssignment, VendorAssignmentHistory
from .serializers import serialize_assignment_history, serialize_staff_profile, serialize_vendor_assignment


def _json_or_form_body(request):
    """Forms bound to `request.POST` only see form-encoded/multipart bodies —
    Django never parses `application/json` into it. The React frontend posts
    plain JSON for these endpoints, so fall back to parsing the raw body
    when POST is empty (keeps working for the rare form-encoded caller too)."""
    if request.POST:
        return request.POST
    try:
        return json.loads(request.body.decode('utf-8') or '{}')
    except json.JSONDecodeError:
        return request.POST


def _sync_overdue_tasks(task_queryset):
    task_queryset.filter(
        due_date__lt=timezone.localdate(),
        task_status__in=[VendorTask.STATUS_PENDING, VendorTask.STATUS_IN_PROGRESS],
    ).update(task_status=VendorTask.STATUS_OVERDUE)


def _log_assignment_audit(user, description, vendor=None, previous_staff=None, new_staff=None):
    log_system_audit(
        user=user,
        action=SystemAuditLog.ACTION_VENDOR_ASSIGNMENT,
        module='vendor_assignment',
        description=description,
        metadata={
            'vendor_id': getattr(vendor, 'vendor_id', ''),
            'vendor_name': getattr(vendor, 'company_name', ''),
            'previous_staff': getattr(previous_staff, 'staff_name', ''),
            'new_staff': getattr(new_staff, 'staff_name', ''),
        },
    )


@transaction.atomic
def _assign_vendor_to_staff(
    *,
    vendor,
    staff_profile,
    actor,
    assignment_role,
    assignment_status,
    start_date,
    end_date,
    reason,
    remarks,
):
    existing_assignment = VendorAssignment.objects.filter(
        vendor=vendor,
        assigned_staff=staff_profile,
        assignment_role=assignment_role,
        assignment_status=VendorAssignment.STATUS_ACTIVE,
    ).first()

    previous_primary = None
    if assignment_role == VendorAssignment.ROLE_PRIMARY and assignment_status == VendorAssignment.STATUS_ACTIVE:
        previous_primary = VendorAssignment.objects.filter(
            vendor=vendor,
            assignment_role=VendorAssignment.ROLE_PRIMARY,
            assignment_status=VendorAssignment.STATUS_ACTIVE,
        ).exclude(assigned_staff=staff_profile).first()

    if previous_primary:
        previous_primary.assignment_status = VendorAssignment.STATUS_REASSIGNED
        previous_primary.end_date = start_date or timezone.localdate()
        previous_primary.remarks = '\n'.join(filter(None, [previous_primary.remarks, remarks]))
        previous_primary.save(update_fields=['assignment_status', 'end_date', 'remarks', 'updated_at'])

    if existing_assignment:
        assignment = existing_assignment
        assignment.start_date = start_date
        assignment.end_date = end_date
        assignment.assignment_status = assignment_status
        assignment.assignment_reason = reason
        assignment.remarks = remarks
        assignment.assigned_by = actor if getattr(actor, 'is_authenticated', False) else None
        assignment.assignment_date = timezone.now()
        assignment.full_clean()
        assignment.save()
    else:
        assignment = VendorAssignment(
            vendor=vendor,
            assigned_staff=staff_profile,
            assigned_by=actor if getattr(actor, 'is_authenticated', False) else None,
            assignment_role=assignment_role,
            assignment_status=assignment_status,
            start_date=start_date,
            end_date=end_date,
            assignment_reason=reason,
            remarks=remarks,
        )
        assignment.full_clean()
        assignment.save()

    VendorAssignmentHistory.objects.create(
        vendor=vendor,
        previous_staff=previous_primary.assigned_staff if previous_primary else None,
        new_staff=staff_profile if assignment_status == VendorAssignment.STATUS_ACTIVE else None,
        changed_by=actor if getattr(actor, 'is_authenticated', False) else None,
        reason=reason,
        remarks=remarks,
    )
    _log_assignment_audit(
        actor,
        f'Assignment updated for vendor {vendor.company_name}.',
        vendor=vendor,
        previous_staff=previous_primary.assigned_staff if previous_primary else None,
        new_staff=staff_profile if assignment_status == VendorAssignment.STATUS_ACTIVE else None,
    )

    if previous_primary:
        log_vendor_activity(
            vendor,
            VendorActivityLog.TYPE_REASSIGNED,
            f'Primary vendor ownership moved from {previous_primary.assigned_staff.staff_name} to {staff_profile.staff_name}.',
            actor,
        )
        create_notification(
            previous_primary.assigned_staff.user,
            'Vendor reassigned',
            f'{vendor.company_name} has been reassigned from your queue.',
            vendor=vendor,
        )
    else:
        log_vendor_activity(
            vendor,
            VendorActivityLog.TYPE_ASSIGNED,
            f'{vendor.company_name} assigned to {staff_profile.staff_name} as {assignment.get_assignment_role_display()}.',
            actor,
        )

    create_notification(
        staff_profile.user,
        'Vendor assigned',
        f'{vendor.company_name} is now assigned to you.',
        vendor=vendor,
    )
    return assignment


def vendor_master(request):
    vendors = Vendor.objects.order_by('company_name')
    query = (request.GET.get('q') or '').strip()
    status_filter = (request.GET.get('status') or '').strip()

    if query:
        vendors = vendors.filter(
            Q(vendor_id__icontains=query)
            | Q(vendor_name__icontains=query)
            | Q(company_name__icontains=query)
            | Q(contact_person__icontains=query)
            | Q(gst_no__icontains=query)
        )
    if status_filter:
        vendors = vendors.filter(status=status_filter)

    if request.method == 'POST':
        if not can_create_vendors(request.user):
            messages.error(request, 'You do not have permission to create vendors.')
            return redirect('procurement-vendor-master')
        form = VendorMasterForm(request.POST)
        if form.is_valid():
            vendor = form.save(commit=False)
            if not vendor.vendor_name:
                vendor.vendor_name = vendor.company_name
            vendor.save()
            messages.success(request, 'Vendor master record created successfully.')
            return redirect('procurement-vendor-master')
    else:
        form = VendorMasterForm(initial={'country': 'India', 'status': 'active'})

    context = {
        'page_title': 'Vendor Master',
        'procurement_nav': True,
        'vendors': vendors,
        'vendor_form': form,
        'query': query,
        'status_filter': status_filter,
    }
    return render(request, 'procurement_vendor_master.html', context)


def vendor_control_dashboard_api(request):
    redirect_response = require_authenticated(request)
    if redirect_response:
        return redirect_response

    vendors_qs = get_accessible_vendor_queryset(request.user)
    tasks_qs = VendorTask.objects.select_related('vendor', 'assigned_staff')
    _sync_overdue_tasks(tasks_qs)

    payload = {
        'is_admin': is_admin_like(request.user),
        'role_label': role_label(request.user),
    }

    if is_admin_like(request.user):
        task_scope = tasks_qs
        active_assignments = VendorAssignment.objects.filter(assignment_status=VendorAssignment.STATUS_ACTIVE)
        payload.update({
            'total_vendors': Vendor.objects.count(),
            'assigned_vendors': active_assignments.filter(assignment_role=VendorAssignment.ROLE_PRIMARY).values('vendor').distinct().count(),
            'unassigned_vendors': Vendor.objects.exclude(
                assignments__assignment_status=VendorAssignment.STATUS_ACTIVE,
                assignments__assignment_role=VendorAssignment.ROLE_PRIMARY,
            ).distinct().count(),
            'pending_tasks': task_scope.filter(task_status__in=[VendorTask.STATUS_PENDING, VendorTask.STATUS_IN_PROGRESS]).count(),
            'overdue_tasks': task_scope.filter(task_status=VendorTask.STATUS_OVERDUE).count(),
            'vendors_by_staff': [
                {'staff_name': row['assigned_staff__staff_name'], 'employee_id': row['assigned_staff__employee_id'], 'total': row['total']}
                for row in active_assignments.values('assigned_staff__staff_name', 'assigned_staff__employee_id')
                .annotate(total=Count('vendor', distinct=True)).order_by('-total', 'assigned_staff__staff_name')
            ],
            'vendors_by_category': [
                {'category': row['vendor_category'] or 'Uncategorized', 'total': row['total']}
                for row in Vendor.objects.values('vendor_category').annotate(total=Count('id')).order_by('-total')
            ],
            'staff_workload': [
                {
                    'staff_name': staff.staff_name,
                    'role': staff.role,
                    'vendor_count': staff.vendor_count,
                    'open_task_count': staff.open_task_count,
                }
                for staff in StaffProfile.objects.filter(is_active=True).annotate(
                    vendor_count=Count(
                        'vendor_assignments',
                        filter=Q(
                            vendor_assignments__assignment_status=VendorAssignment.STATUS_ACTIVE,
                            vendor_assignments__assignment_role=VendorAssignment.ROLE_PRIMARY,
                        ),
                        distinct=True,
                    ),
                    open_task_count=Count(
                        'assigned_tasks',
                        filter=Q(
                            assigned_tasks__task_status__in=[
                                VendorTask.STATUS_PENDING,
                                VendorTask.STATUS_IN_PROGRESS,
                                VendorTask.STATUS_OVERDUE,
                            ]
                        ),
                        distinct=True,
                    ),
                ).order_by('-vendor_count', '-open_task_count', 'staff_name')
            ],
            'recently_reassigned': [
                serialize_assignment_history(row)
                for row in VendorAssignmentHistory.objects.select_related('vendor', 'previous_staff', 'new_staff', 'changed_by')
                .exclude(previous_staff=None).order_by('-changed_date')[:8]
            ],
        })
    else:
        profile = get_staff_profile(request.user)
        task_scope = tasks_qs.filter(vendor__in=vendors_qs) if profile else tasks_qs.none()
        upcoming_cutoff = timezone.localdate() + timedelta(days=7)
        payload.update({
            'my_vendor_count': vendors_qs.count(),
            'pending_tasks': task_scope.filter(task_status__in=[VendorTask.STATUS_PENDING, VendorTask.STATUS_IN_PROGRESS]).count(),
            'overdue_tasks': task_scope.filter(task_status=VendorTask.STATUS_OVERDUE).count(),
            'completed_tasks': task_scope.filter(task_status=VendorTask.STATUS_COMPLETED).count(),
            'upcoming_followups': [
                {
                    'id': task.id,
                    'task_title': task.task_title,
                    'vendor_name': task.vendor.company_name,
                    'due_date': task.due_date.isoformat(),
                    'priority': task.priority,
                    'task_status': task.task_status,
                }
                for task in task_scope.filter(
                    due_date__gte=timezone.localdate(),
                    due_date__lte=upcoming_cutoff,
                    task_status__in=[VendorTask.STATUS_PENDING, VendorTask.STATUS_IN_PROGRESS, VendorTask.STATUS_OVERDUE],
                ).select_related('vendor').order_by('due_date')[:8]
            ],
            'vendor_documents_pending': vendors_qs.filter(
                Q(passbook_file__isnull=True) | Q(passbook_file='') | Q(pan_no='') | Q(gst_no='')
            ).count(),
            'my_assignments': [
                {'vendor_id': a.vendor.vendor_id, 'vendor_name': a.vendor.company_name, 'assignment_role': a.assignment_role}
                for a in get_active_vendor_assignments(request.user).select_related('vendor')[:8]
            ],
        })

    return JsonResponse(payload)


def staff_list_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response

    query = (request.GET.get('q') or '').strip()
    role_filter = (request.GET.get('role') or '').strip()

    staff_rows = StaffProfile.objects.select_related('user', 'reporting_manager').order_by('staff_name')
    if query:
        staff_rows = staff_rows.filter(
            Q(staff_name__icontains=query)
            | Q(employee_id__icontains=query)
            | Q(department__icontains=query)
            | Q(user__username__icontains=query)
        )
    if role_filter:
        staff_rows = staff_rows.filter(role=role_filter)

    return JsonResponse({
        'results': [serialize_staff_profile(staff) for staff in staff_rows],
        'role_choices': StaffProfile._meta.get_field('role').choices,
    })


def staff_create_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    form = StaffProfileForm(_json_or_form_body(request))
    if not form.is_valid():
        return JsonResponse({'error': 'Please correct the errors and try again.', 'field_errors': form.errors}, status=400)
    staff = form.save()
    return JsonResponse({'message': 'Staff profile saved successfully.', 'staff': serialize_staff_profile(staff)}, status=201)


def available_users_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response

    from django.contrib.auth import get_user_model
    User = get_user_model()
    rows = [
        {'id': user.id, 'username': user.username, 'full_name': user.get_full_name()}
        for user in User.objects.order_by('username')
    ]
    return JsonResponse({'results': rows})


def vendor_assignment_create_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    staff_queryset = StaffProfile.objects.filter(is_active=True).order_by('staff_name')
    form = VendorAssignmentForm(_json_or_form_body(request))
    form.fields['assigned_staff'].queryset = staff_queryset
    form.fields['vendor'].queryset = Vendor.objects.order_by('company_name')
    if not form.is_valid():
        return JsonResponse({'error': 'Please correct the errors and try again.', 'field_errors': form.errors}, status=400)

    assignment = form.save(commit=False)
    assignment = _assign_vendor_to_staff(
        vendor=assignment.vendor,
        staff_profile=assignment.assigned_staff,
        actor=request.user,
        assignment_role=assignment.assignment_role,
        assignment_status=assignment.assignment_status,
        start_date=assignment.start_date,
        end_date=assignment.end_date,
        reason=assignment.assignment_reason,
        remarks=assignment.remarks,
    )
    return JsonResponse({'message': 'Vendor assignment saved.', 'assignment': serialize_vendor_assignment(assignment)}, status=201)


def vendor_assignment_remove_api(request, assignment_id):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    body = _json_or_form_body(request)
    assignment = get_object_or_404(VendorAssignment, pk=assignment_id)
    assignment.assignment_status = VendorAssignment.STATUS_REMOVED
    assignment.end_date = timezone.localdate()
    assignment.save(update_fields=['assignment_status', 'end_date', 'updated_at'])
    VendorAssignmentHistory.objects.create(
        vendor=assignment.vendor,
        previous_staff=assignment.assigned_staff,
        new_staff=None,
        changed_by=request.user,
        reason='Assignment removed',
        remarks=(body.get('remarks') or ''),
    )
    _log_assignment_audit(
        request.user,
        f'Assignment removed for vendor {assignment.vendor.company_name}.',
        vendor=assignment.vendor,
        previous_staff=assignment.assigned_staff,
        new_staff=None,
    )
    log_vendor_activity(
        assignment.vendor,
        VendorActivityLog.TYPE_REASSIGNED,
        f'Assignment removed from {assignment.assigned_staff.staff_name}.',
        request.user,
    )
    return JsonResponse({'message': 'Vendor assignment removed.', 'assignment': serialize_vendor_assignment(assignment)})


def vendor_bulk_assign_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    staff_queryset = StaffProfile.objects.filter(is_active=True).order_by('staff_name')
    bulk_form = VendorBulkAssignForm(request.POST, staff_queryset=staff_queryset)
    if not bulk_form.is_valid():
        return JsonResponse({'error': 'Please correct the errors and try again.', 'field_errors': bulk_form.errors}, status=400)

    selected_staff = bulk_form.cleaned_data['assigned_staff']
    selected_vendors = bulk_form.cleaned_data['vendors']
    for vendor in selected_vendors:
        _assign_vendor_to_staff(
            vendor=vendor,
            staff_profile=selected_staff,
            actor=request.user,
            assignment_role=bulk_form.cleaned_data['assignment_role'],
            assignment_status=VendorAssignment.STATUS_ACTIVE,
            start_date=timezone.localdate(),
            end_date=None,
            reason=bulk_form.cleaned_data['assignment_reason'],
            remarks=bulk_form.cleaned_data['remarks'],
        )
    return JsonResponse({'message': f'{selected_vendors.count()} vendors assigned to {selected_staff.staff_name}.'})


def vendor_bulk_assign_upload_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    upload_form = VendorAssignmentUploadForm(request.POST, request.FILES)
    if not upload_form.is_valid():
        return JsonResponse({'error': 'Please correct the errors and try again.', 'field_errors': upload_form.errors}, status=400)

    try:
        from openpyxl import load_workbook
    except ImportError:
        bundled_path = Path.home() / '.cache' / 'codex-runtimes' / 'codex-primary-runtime' / 'dependencies' / 'python'
        if bundled_path.exists():
            sys.path.append(str(bundled_path))
            from openpyxl import load_workbook
        else:
            return JsonResponse({'error': 'Excel upload support is unavailable until openpyxl is installed.'}, status=400)

    workbook = load_workbook(filename=BytesIO(upload_form.cleaned_data['assignment_file'].read()), data_only=True)
    sheet = workbook.active
    header = [str(cell.value).strip().lower() if cell.value is not None else '' for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    vendor_idx = header.index('vendor_id') if 'vendor_id' in header else -1
    employee_idx = header.index('employee_id') if 'employee_id' in header else -1
    if vendor_idx < 0 or employee_idx < 0:
        return JsonResponse({'error': 'Excel must include vendor_id and employee_id columns.'}, status=400)

    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        vendor_id = str(row[vendor_idx]).strip() if row[vendor_idx] else ''
        employee_id = str(row[employee_idx]).strip() if row[employee_idx] else ''
        if not vendor_id or not employee_id:
            continue
        try:
            vendor = Vendor.objects.get(vendor_id=vendor_id)
            staff_profile = StaffProfile.objects.get(employee_id=employee_id, is_active=True)
        except (Vendor.DoesNotExist, StaffProfile.DoesNotExist):
            continue
        _assign_vendor_to_staff(
            vendor=vendor,
            staff_profile=staff_profile,
            actor=request.user,
            assignment_role=upload_form.cleaned_data['default_role'],
            assignment_status=VendorAssignment.STATUS_ACTIVE,
            start_date=timezone.localdate(),
            end_date=None,
            reason='Bulk Excel assignment',
            remarks='Uploaded from assignment workbook.',
        )
        count += 1
    return JsonResponse({'message': f'{count} vendor assignments processed from Excel.'})


def vendor_auto_distribute_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    staff_queryset = StaffProfile.objects.filter(is_active=True).order_by('staff_name')
    auto_form = VendorAutoDistributeForm(request.POST, staff_queryset=staff_queryset)
    if not auto_form.is_valid():
        return JsonResponse({'error': 'Please correct the errors and try again.', 'field_errors': auto_form.errors}, status=400)

    selected_staff = list(auto_form.cleaned_data['staff_members'])
    if not selected_staff:
        return JsonResponse({'error': 'Select at least one staff member to distribute vendors to.'}, status=400)

    selected_vendors = list(auto_form.cleaned_data['vendors'])
    if not selected_vendors:
        selected_vendors = list(
            Vendor.objects.exclude(
                assignments__assignment_status=VendorAssignment.STATUS_ACTIVE,
                assignments__assignment_role=VendorAssignment.ROLE_PRIMARY,
            ).order_by('company_name')
        )

    if auto_form.cleaned_data['strategy'] == VendorAutoDistributeForm.STRATEGY_LOCATION:
        selected_vendors.sort(key=lambda vendor: (vendor.state or '', vendor.city or '', vendor.company_name))
    elif auto_form.cleaned_data['strategy'] == VendorAutoDistributeForm.STRATEGY_CATEGORY:
        selected_vendors.sort(key=lambda vendor: (vendor.vendor_category or '', vendor.company_name))
    else:
        staff_load = {
            staff.id: VendorAssignment.objects.filter(
                assigned_staff=staff,
                assignment_status=VendorAssignment.STATUS_ACTIVE,
            ).count()
            for staff in selected_staff
        }
        selected_staff.sort(key=lambda staff: (staff_load[staff.id], staff.staff_name))

    for index, vendor in enumerate(selected_vendors):
        target_staff = selected_staff[index % len(selected_staff)]
        _assign_vendor_to_staff(
            vendor=vendor,
            staff_profile=target_staff,
            actor=request.user,
            assignment_role=auto_form.cleaned_data['assignment_role'],
            assignment_status=VendorAssignment.STATUS_ACTIVE,
            start_date=timezone.localdate(),
            end_date=None,
            reason=f'Auto distributed by {auto_form.cleaned_data["strategy"]}.',
            remarks='Bulk auto distribution.',
        )
    return JsonResponse({'message': f'{len(selected_vendors)} vendors auto-distributed successfully.'})


def vendor_distribution_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response

    active_assignments = VendorAssignment.objects.filter(
        assignment_status=VendorAssignment.STATUS_ACTIVE
    ).select_related('vendor', 'assigned_staff')
    distribution_map = defaultdict(list)
    for assignment in active_assignments:
        distribution_map[assignment.assigned_staff].append(assignment)

    rows = [
        {
            'staff_name': staff.staff_name,
            'role': staff.role,
            'primary_count': sum(1 for row in assignments if row.assignment_role == VendorAssignment.ROLE_PRIMARY),
            'support_count': sum(1 for row in assignments if row.assignment_role == VendorAssignment.ROLE_SUPPORTING),
            'assignments': [
                {'vendor_id': row.vendor.vendor_id, 'vendor_name': row.vendor.company_name, 'assignment_role': row.assignment_role}
                for row in sorted(assignments, key=lambda row: (row.assignment_role, row.vendor.company_name))
            ],
        }
        for staff, assignments in distribution_map.items()
    ]
    rows.sort(key=lambda row: (-row['primary_count'], -row['support_count'], row['staff_name']))

    return JsonResponse({'results': rows})


def staff_performance_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response

    _sync_overdue_tasks(VendorTask.objects.all())
    performance_rows = StaffProfile.objects.filter(is_active=True).annotate(
        active_vendor_count=Count(
            'vendor_assignments',
            filter=Q(vendor_assignments__assignment_status=VendorAssignment.STATUS_ACTIVE),
            distinct=True,
        ),
        pending_task_count=Count(
            'assigned_tasks',
            filter=Q(assigned_tasks__task_status__in=[VendorTask.STATUS_PENDING, VendorTask.STATUS_IN_PROGRESS]),
            distinct=True,
        ),
        overdue_task_count=Count(
            'assigned_tasks',
            filter=Q(assigned_tasks__task_status=VendorTask.STATUS_OVERDUE),
            distinct=True,
        ),
        completed_task_count=Count(
            'assigned_tasks',
            filter=Q(assigned_tasks__task_status=VendorTask.STATUS_COMPLETED),
            distinct=True,
        ),
    ).order_by('-active_vendor_count', '-pending_task_count', 'staff_name')

    rows = [
        {
            'staff_name': staff.staff_name,
            'role': staff.role,
            'active_vendor_count': staff.active_vendor_count,
            'pending_task_count': staff.pending_task_count,
            'overdue_task_count': staff.overdue_task_count,
            'completed_task_count': staff.completed_task_count,
        }
        for staff in performance_rows
    ]
    return JsonResponse({'results': rows})


def my_vendors_api(request):
    redirect_response = require_authenticated(request)
    if redirect_response:
        return redirect_response

    vendors_qs = get_accessible_vendor_queryset(request.user).order_by('company_name')
    query = (request.GET.get('q') or '').strip()
    if query:
        vendors_qs = vendors_qs.filter(
            Q(vendor_id__icontains=query)
            | Q(company_name__icontains=query)
            | Q(city__icontains=query)
            | Q(vendor_category__icontains=query)
        )

    vendor_ids = list(vendors_qs.values_list('id', flat=True))
    task_counts = {
        row['vendor_id']: row['total']
        for row in VendorTask.objects.filter(
            vendor_id__in=vendor_ids,
            task_status__in=[VendorTask.STATUS_PENDING, VendorTask.STATUS_IN_PROGRESS, VendorTask.STATUS_OVERDUE],
        ).values('vendor_id').annotate(total=Count('id'))
    }
    primary_assignments = {
        row.vendor_id: row.assigned_staff.staff_name
        for row in VendorAssignment.objects.filter(
            vendor_id__in=vendor_ids,
            assignment_status=VendorAssignment.STATUS_ACTIVE,
            assignment_role=VendorAssignment.ROLE_PRIMARY,
        ).select_related('assigned_staff')
    }

    rows = [
        {
            'vendor_id': vendor.vendor_id,
            'company_name': vendor.company_name,
            'category': vendor.vendor_category,
            'city': vendor.city,
            'state': vendor.state,
            'task_count': task_counts.get(vendor.id, 0),
            'primary_staff': primary_assignments.get(vendor.id, ''),
        }
        for vendor in vendors_qs
    ]

    return JsonResponse({'results': rows, 'is_admin': is_admin_like(request.user)})


def vendor_control_detail_api(request, vendor_id):
    redirect_response = require_authenticated(request)
    if redirect_response:
        return redirect_response

    vendor = get_object_or_404(Vendor, vendor_id=vendor_id)
    try:
        ensure_vendor_access(request.user, vendor)
    except PermissionDenied as exc:
        return JsonResponse({'error': str(exc) or 'You are not authorized to access this vendor.'}, status=403)

    assignments = vendor.assignments.select_related('assigned_staff', 'assigned_by').order_by('-updated_at')
    tasks = vendor.vendor_tasks.select_related('assigned_staff', 'created_by').order_by('due_date', '-created_at')

    payload = {
        'vendor': {
            'vendor_id': vendor.vendor_id,
            'company_name': vendor.company_name,
            'vendor_name': vendor.vendor_name,
            'vendor_category': vendor.vendor_category,
            'gst_no': vendor.gst_no,
            'pan_no': vendor.pan_no,
            'contact_person': vendor.contact_person,
            'mobile_number': vendor.mobile_number,
            'email_id': vendor.email_id,
            'address': vendor.address,
            'city': vendor.city,
            'state': vendor.state,
            'pin_code': vendor.pin_code,
            'country': vendor.country,
            'bank_details': vendor.bank_details,
            'status': vendor.status,
        },
        'assignments': [serialize_vendor_assignment(a) for a in assignments],
        'tasks': [
            {
                'id': t.id,
                'task_title': t.task_title,
                'assigned_staff': t.assigned_staff.staff_name,
                'due_date': t.due_date.isoformat(),
                'task_status': t.task_status,
                'priority': t.priority,
            }
            for t in tasks
        ],
        'activity': [
            {
                'id': a.id,
                'activity_type': a.activity_type,
                'activity_type_display': a.get_activity_type_display(),
                'description': a.description,
                'performed_by': a.performed_by.username if a.performed_by_id else 'System',
                'created_at': a.created_at.isoformat(),
            }
            for a in vendor.activity_logs.select_related('performed_by').all()[:25]
        ],
        'purchase_orders': [
            {
                'id': po.id,
                'po_number': po.po_number,
                'po_date': po.po_date.isoformat() if po.po_date else '',
                'total_po_value': str(po.total_po_value or 0),
                'status': po.status,
            }
            for po in PurchaseOrder.objects.filter(vendor=vendor).order_by('-po_date')[:10]
        ],
        'can_edit_vendor_notes': not is_view_only(request.user),
    }
    return JsonResponse(payload)


def vendor_note_create_api(request, vendor_id):
    redirect_response = require_authenticated(request)
    if redirect_response:
        return redirect_response
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    vendor = get_object_or_404(Vendor, vendor_id=vendor_id)
    try:
        ensure_vendor_write_access(request.user, vendor)
    except PermissionDenied as exc:
        return JsonResponse({'error': str(exc) or 'You are not authorized to modify this vendor.'}, status=403)

    note_form = VendorNoteForm(_json_or_form_body(request))
    if not note_form.is_valid():
        return JsonResponse({'error': 'Please correct the errors and try again.', 'field_errors': note_form.errors}, status=400)

    entry = log_vendor_activity(
        vendor,
        VendorActivityLog.TYPE_NOTE_ADDED,
        note_form.cleaned_data['note'],
        request.user,
    )
    return JsonResponse({
        'message': 'Vendor note added.',
        'activity': {
            'id': entry.id,
            'activity_type': entry.activity_type,
            'activity_type_display': entry.get_activity_type_display(),
            'description': entry.description,
            'performed_by': entry.performed_by.username if entry.performed_by_id else 'System',
            'created_at': entry.created_at.isoformat(),
        },
    }, status=201)


def accessible_vendors_api(request):
    redirect_response = require_authenticated(request)
    if redirect_response:
        return redirect_response

    vendor_rows = [
        {
            'id': vendor.id,
            'vendor_id': vendor.vendor_id,
            'company_name': vendor.company_name,
            'category': vendor.vendor_category,
            'city': vendor.city,
            'state': vendor.state,
        }
        for vendor in get_accessible_vendor_queryset(request.user).order_by('company_name')
    ]
    return JsonResponse({'results': vendor_rows})


def assignment_history_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response

    query = (request.GET.get('q') or '').strip()
    history_rows = VendorAssignmentHistory.objects.select_related('vendor', 'previous_staff', 'new_staff', 'changed_by')
    if query:
        history_rows = history_rows.filter(
            Q(vendor__vendor_id__icontains=query)
            | Q(vendor__company_name__icontains=query)
            | Q(previous_staff__staff_name__icontains=query)
            | Q(new_staff__staff_name__icontains=query)
        )

    rows = [serialize_assignment_history(row) for row in history_rows.order_by('-changed_date')[:100]]
    return JsonResponse({'results': rows})


def assignments_api(request):
    redirect_response = require_admin_access(request)
    if redirect_response:
        return redirect_response

    query = (request.GET.get('q') or '').strip()
    status_filter = (request.GET.get('status') or '').strip()
    assignment_rows = VendorAssignment.objects.select_related('vendor', 'assigned_staff', 'assigned_by').order_by('-updated_at')
    if query:
        assignment_rows = assignment_rows.filter(
            Q(vendor__vendor_id__icontains=query)
            | Q(vendor__company_name__icontains=query)
            | Q(assigned_staff__staff_name__icontains=query)
        )
    if status_filter:
        assignment_rows = assignment_rows.filter(assignment_status=status_filter)

    rows = [serialize_vendor_assignment(row) for row in assignment_rows[:200]]
    return JsonResponse({'results': rows, 'status_choices': VendorAssignment.STATUS_CHOICES})
